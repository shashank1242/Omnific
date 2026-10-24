"""Document parsers. Each returns an ordered list of Blocks (heading | text | table | figure)."""

import csv
import io
import os
import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from typing import Literal

BlockKind = Literal["heading", "text", "table", "figure"]


class UnsupportedFile(ValueError):
    """The file type is not supported (retrying cannot help)."""


class CorruptFile(ValueError):
    """The file could not be parsed (retrying cannot help)."""


@dataclass
class Block:
    kind: BlockKind
    page: int | None = None
    section: str | None = None
    text: str | None = None
    image: bytes | None = None
    mime: str | None = None


@dataclass
class Parsed:
    blocks: list[Block] = field(default_factory=list)
    pages: int = 1
    title: str | None = None


MAX_TABLE_ROWS = int(os.environ.get("RAG_MAX_TABLE_ROWS", "200"))
RENDER_DPI = 150
_EXT = {
    ".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv",
    ".html": "html", ".htm": "html", ".md": "md", ".markdown": "md", ".txt": "txt",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image",
}
_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/csv": "csv", "text/html": "html", "text/markdown": "md", "text/plain": "txt",
    "image/png": "image", "image/jpeg": "image", "image/webp": "image",
}


def detect_format(filename: str, content_type: str = "") -> str:
    ext = os.path.splitext(filename.lower())[1]
    fmt = _EXT.get(ext) or _MIME.get((content_type or "").split(";")[0].strip().lower())
    if not fmt:
        raise UnsupportedFile(f"unsupported file type: {filename!r} ({content_type})")
    return fmt


def parse(data: bytes, filename: str, content_type: str = "") -> Parsed:
    fmt = detect_format(filename, content_type)
    try:
        result = _PARSERS[fmt](data)
    except (UnsupportedFile, CorruptFile):
        raise
    except Exception as e:
        raise CorruptFile(f"could not parse {fmt} file: {type(e).__name__}: {e}") from e
    return result


# --- helpers ----------------------------------------------------------------------------------

def _cell(v: object) -> str:
    s = "" if v is None else str(v)
    return re.sub(r"\s*\n\s*", " ", s).replace("|", "\\|").strip()


def markdown_table(rows: list[list[object]], max_rows: int | None = None) -> str:
    """Rows (first row = header) to a Markdown table; ragged rows are padded."""
    rows = [r for r in rows if any(_cell(c) for c in r)]
    if not rows:
        return ""
    cap = MAX_TABLE_ROWS if max_rows is None else max_rows
    truncated = len(rows) - 1 > cap
    if truncated:
        rows = rows[: cap + 1]
    width = max(len(r) for r in rows)
    norm = [[_cell(c) for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "|".join([" --- "] * width) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    if truncated:
        lines.append(f"\n(table truncated to {cap} rows)")
    return "\n".join(lines)


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


class _Builder:
    """Collects blocks and tracks the current section heading."""

    def __init__(self) -> None:
        self.blocks: list[Block] = []
        self.section: str | None = None
        self.title: str | None = None

    def heading(self, text: str, page: int | None = None) -> None:
        text = " ".join(text.split())
        if not text:
            return
        self.section = text
        self.blocks.append(Block("heading", page, text, text))

    def text(self, text: str, page: int | None = None) -> None:
        if text.strip():
            self.blocks.append(Block("text", page, self.section, text.strip()))

    def table(self, rows: list[list[object]], page: int | None = None) -> None:
        md = markdown_table(rows)
        if md:
            self.blocks.append(Block("table", page, self.section, md))

    def figure(self, image: bytes, mime: str, page: int | None = None) -> None:
        self.blocks.append(Block("figure", page, self.section, None, image, mime))


# --- PDF --------------------------------------------------------------------------------------

def _is_bold(fontname: str) -> bool:
    f = fontname.lower()
    return any(k in f for k in ("bold", "black", "heavy", "semibold", "demi")) or f.endswith(",b")


def _inside(obj: dict, bboxes: list[tuple]) -> bool:
    cx, cy = (obj["x0"] + obj["x1"]) / 2, (obj["top"] + obj["bottom"]) / 2
    return any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in bboxes)


def _pdf_lines(words: list[dict]) -> list[dict]:
    words = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines: list[list[dict]] = []
    for w in words:
        if lines and abs(w["top"] - lines[-1][0]["top"]) < 0.5 * w["size"]:
            lines[-1].append(w)
        else:
            lines.append([w])
    out = []
    for ws in lines:
        ws.sort(key=lambda w: w["x0"])
        sizes = [w["size"] for w in ws]
        out.append({
            "text": " ".join(w["text"] for w in ws),
            "size": statistics.median(sizes),
            "bold": all(_is_bold(w["fontname"]) for w in ws),
            "top": min(w["top"] for w in ws), "bottom": max(w["bottom"] for w in ws),
            "n": sum(len(w["text"]) for w in ws),
        })
    return out


def _parse_pdf(data: bytes) -> Parsed:
    import pdfplumber
    import pypdfium2

    b = _Builder()
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n_pages = len(pdf.pages)
        if n_pages == 0:
            raise CorruptFile("PDF has no pages")
        per_page = []
        sizes: Counter = Counter()
        for page in pdf.pages:
            tables = page.find_tables()
            tboxes = [t.bbox for t in tables]
            words = [w for w in page.extract_words(extra_attrs=["size", "fontname"])
                     if not _inside(w, tboxes)]
            for w in words:
                sizes[round(w["size"] * 2) / 2] += len(w["text"])
            per_page.append((page, tables, tboxes, words))
        body = sizes.most_common(1)[0][0] if sizes else 10.0

        pdfium_doc = None
        for pno, (page, tables, tboxes, words) in enumerate(per_page, start=1):
            items: list[tuple[float, str, object]] = [(t.bbox[1], "table", t) for t in tables]
            items += [(ln["top"], "line", ln) for ln in _pdf_lines(words)]
            items.sort(key=lambda it: it[0])
            para: list[str] = []
            prev = None
            head_prev = None

            def flush(para: list[str] = para, pno: int = pno) -> None:
                if para:
                    b.text(" ".join(para), pno)
                    para.clear()

            for _, kind, obj in items:
                if kind == "table":
                    flush()
                    head_prev = None
                    b.table(obj.extract(), pno)  # type: ignore[attr-defined]
                    prev = None
                    continue
                ln = obj  # type: ignore[assignment]
                text = ln["text"]
                heading = len(text) <= 120 and (
                    ln["size"] >= body * 1.15 or (ln["bold"] and ln["size"] >= body * 0.95
                                                  and not text.rstrip().endswith((".", ",", ";")))
                )
                if heading:
                    flush()
                    if (head_prev and abs(head_prev["size"] - ln["size"]) < 0.6
                            and ln["top"] - head_prev["bottom"] < 0.6 * ln["size"]
                            and b.blocks and b.blocks[-1].kind == "heading"):
                        merged = b.blocks[-1].text + " " + text  # type: ignore[operator]
                        b.blocks.pop()
                        b.heading(merged, pno)
                    else:
                        b.heading(text, pno)
                    head_prev = ln
                    prev = None
                    continue
                head_prev = None
                if prev is not None and ln["top"] - prev["bottom"] > 0.5 * ln["size"]:
                    flush()
                para.append(text)
                prev = ln
            flush()

            area = max(1.0, page.width * page.height)
            big_images = [im for im in page.images
                          if (im["x1"] - im["x0"]) * (im["bottom"] - im["top"]) >= 0.04 * area]
            vectors = [o for o in (*page.curves, *page.lines, *page.rects) if not _inside(o, tboxes)]
            if big_images or len(vectors) > 25:
                if pdfium_doc is None:
                    pdfium_doc = pypdfium2.PdfDocument(data)
                pil = pdfium_doc[pno - 1].render(scale=RENDER_DPI / 72).to_pil().convert("RGB")
                buf = io.BytesIO()
                pil.save(buf, "PNG")
                b.figure(buf.getvalue(), "image/png", pno)
        meta_title = (pdf.metadata or {}).get("Title")
        if pdfium_doc is not None:
            pdfium_doc.close()
    title = meta_title.strip() if isinstance(meta_title, str) and meta_title.strip() else None
    if title is None:
        title = next((bl.text for bl in b.blocks if bl.kind == "heading"), None)
    return Parsed(b.blocks, n_pages, title)


# --- DOCX -------------------------------------------------------------------------------------

def _parse_docx(data: bytes) -> Parsed:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    b = _Builder()
    ns_p, ns_tbl = (d.element.body.tag.rsplit("}", 1)[0] + "}" + t for t in ("p", "tbl"))
    for el in d.element.body.iterchildren():
        if el.tag == ns_tbl:
            b.table([[c.text for c in row.cells] for row in Table(el, d).rows])
        elif el.tag == ns_p:
            p = Paragraph(el, d)
            style = (p.style.name if p.style is not None else "") or ""
            text = p.text.strip()
            if text and (style.startswith("Heading") or style == "Title"):
                b.heading(text)
                if style == "Title" and b.title is None:
                    b.title = text
            elif text:
                b.text(text)
            for rid in el.xpath(".//a:blip/@r:embed"):
                part = d.part.related_parts.get(rid)
                if part is not None and part.content_type.startswith("image/"):
                    b.figure(part.blob, part.content_type)
    title = b.title or (d.core_properties.title or "").strip() or None
    if title is None:
        title = next((bl.text for bl in b.blocks if bl.kind == "heading"), None)
    return Parsed(b.blocks, 1, title)


# --- PPTX -------------------------------------------------------------------------------------

def _pptx_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # GROUP
            yield from _pptx_shapes(sh.shapes)
        else:
            yield sh


def _parse_pptx(data: bytes) -> Parsed:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    b = _Builder()
    n = 0
    for n, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
            b.heading(title_shape.text_frame.text, n)
        else:
            b.section = f"Slide {n}"
        for sh in _pptx_shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                b.table([[c.text for c in row.cells] for row in sh.table.rows], n)
            elif sh.shape_type == 13:  # PICTURE
                b.figure(sh.image.blob, sh.image.content_type, n)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                b.text("\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip()), n)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                b.text("Speaker notes: " + notes, n)
    title = (prs.core_properties.title or "").strip() or None
    if title is None:
        title = next((bl.text for bl in b.blocks if bl.kind == "heading"), None)
    return Parsed(b.blocks, max(n, 1), title)


# --- XLSX / CSV -------------------------------------------------------------------------------

def _parse_xlsx(data: bytes) -> Parsed:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    b = _Builder()
    for i, ws in enumerate(wb.worksheets, start=1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_TABLE_ROWS + 1:
                break
        b.heading(f"Sheet: {ws.title}", i)
        b.table(rows, i)
    wb.close()
    title = ((wb.properties.title or "").strip() or None) if wb.properties else None
    return Parsed(b.blocks, max(len(wb.worksheets), 1), title)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1")


def _parse_csv(data: bytes) -> Parsed:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append(row)
        if len(rows) > MAX_TABLE_ROWS + 1:
            break
    b = _Builder()
    b.table(rows, 1)
    return Parsed(b.blocks, 1, None)


# --- HTML -------------------------------------------------------------------------------------

def _parse_html(data: bytes) -> Parsed:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    for t in soup(["script", "style", "noscript", "template"]):
        t.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else None
    b = _Builder()
    root = soup.body or soup
    for el in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote", "table"]):
        if el.name != "table" and el.find_parent("table"):
            continue
        if el.name in ("p", "blockquote") and el.find_parent(["li", "blockquote"]):
            continue
        if el.name[0] == "h" and len(el.name) == 2:
            b.heading(el.get_text(" ", strip=True))
        elif el.name == "table":
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
                    for tr in el.find_all("tr")]
            b.table(rows)
        elif el.name == "pre":
            b.text(el.get_text())
        else:
            b.text(el.get_text(" ", strip=True))
    if title is None:
        title = next((bl.text for bl in b.blocks if bl.kind == "heading"), None)
    return Parsed(b.blocks, 1, title or None)


# --- Markdown / text --------------------------------------------------------------------------

_MD_HEAD = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")


def _parse_md(data: bytes) -> Parsed:
    b = _Builder()
    para: list[str] = []
    table: list[str] = []
    in_code = False

    def flush_para() -> None:
        if para:
            b.text("\n".join(para))
            para.clear()

    def flush_table() -> None:
        if table:
            b.blocks.append(Block("table", None, b.section, "\n".join(table)))
            table.clear()

    for line in _decode(data).splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            para.append(line)
            continue
        if in_code:
            para.append(line)
            continue
        if line.lstrip().startswith("|"):
            flush_para()
            table.append(line.strip())
            continue
        flush_table()
        m = _MD_HEAD.match(line)
        if m:
            flush_para()
            b.heading(m.group(2))
        elif not line.strip():
            flush_para()
        else:
            para.append(line)
    flush_para()
    flush_table()
    title = next((bl.text for bl in b.blocks if bl.kind == "heading"), None)
    return Parsed(b.blocks, 1, title)


def _parse_txt(data: bytes) -> Parsed:
    b = _Builder()
    for p in _paragraphs(_decode(data)):
        b.text(p)
    return Parsed(b.blocks, 1, None)


# --- images -----------------------------------------------------------------------------------

def _parse_image(data: bytes) -> Parsed:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as im:
        im.load()
        fmt = (im.format or "").upper()
    mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}.get(fmt)
    if mime is None:
        raise UnsupportedFile(f"unsupported image format {fmt!r}")
    b = _Builder()
    b.figure(data, mime, 1)
    return Parsed(b.blocks, 1, None)


_PARSERS = {
    "pdf": _parse_pdf, "docx": _parse_docx, "pptx": _parse_pptx, "xlsx": _parse_xlsx,
    "csv": _parse_csv, "html": _parse_html, "md": _parse_md, "txt": _parse_txt, "image": _parse_image,
}
