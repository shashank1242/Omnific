"""Programmatically built fixture files (no binary blobs in the repo)."""

import csv
import io

BODY1 = ["The wind tunnel test measured drag on a scale model of the vehicle.",
         "Each run lasted sixty seconds at a fixed free stream speed."]
BODY2 = ["Models were mounted on a force balance with a sting support."]


def _text(x, y, font, size, s):
    s = s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"BT /{font} {size} Tf {x} {y} Td ({s}) Tj ET\n"


def make_pdf() -> bytes:
    p1 = _text(72, 720, "F2", 20, "Wind Tunnel Report")
    for i, line in enumerate(BODY1):
        p1 += _text(72, 690 - i * 14, "F1", 10, line)
    p1 += _text(72, 640, "F2", 14, "2. Method")
    for i, line in enumerate(BODY2):
        p1 += _text(72, 618 - i * 14, "F1", 10, line)
    # 3x3 table: grid lines + cell text
    xs, ys = [72, 172, 272, 372], [560, 540, 520, 500]
    for y in ys:
        p1 += f"{xs[0]} {y} m {xs[-1]} {y} l S\n"
    for x in xs:
        p1 += f"{x} {ys[0]} m {x} {ys[-1]} l S\n"
    cells = [["Speed", "Drag", "Unit"], ["10", "1.5", "N"], ["20", "6.1", "N"]]
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            p1 += _text(xs[c] + 5, ys[r] - 14, "F1", 10, val)
    # page 2: heading + a vector chart (many line segments)
    p2 = _text(72, 720, "F2", 14, "3. Results")
    p2 += _text(72, 700, "F1", 10, "Drag rises with speed as shown in the figure below.")
    p2 += "72 300 m 472 300 l S\n72 300 m 72 600 l S\n"
    for i in range(30):
        x0, x1 = 80 + i * 12, 92 + i * 12
        p2 += f"{x0} {310 + i * i // 4} m {x1} {310 + (i + 1) * (i + 1) // 4} l S\n"
    streams = [p1, p2]

    objs: list[bytes] = []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(b"<< /Type /Pages /Kids [5 0 R 7 0 R] /Count 2 >>")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")
    res = b"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >>"
    for i, s in enumerate(streams):
        page_no = 5 + 2 * i
        objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] " + res
                    + f" /Contents {page_no + 1} 0 R >>".encode())
        data = s.encode("latin-1")
        objs.append(f"<< /Length {len(data)} >>\nstream\n".encode() + data + b"endstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_png(color=(200, 30, 30), size=(64, 48)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Hull Design Notes", level=0)
    d.add_heading("Materials", level=1)
    d.add_paragraph("The hull is laminated carbon fibre over a foam core.")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Part", "Mass kg"), ("Hull", "12.5"), ("Keel", "40")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_heading("Loads", level=1)
    d.add_paragraph("Peak slamming load is 3 g.")
    d.add_picture(io.BytesIO(make_png()))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Launch Overview"
    s1.placeholders[1].text = "Thrust is 50 kN at sea level."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Budget"
    gt = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), v in {(0, 0): "Item", (0, 1): "Cost", (1, 0): "Engine", (1, 1): "5"}.items():
        gt.cell(r, c).text = v
    s2.shapes.add_picture(io.BytesIO(make_png()), Inches(1), Inches(4))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(["Name", "Density"])
    ws.append(["Steel", 7850])
    ws.append(["Water", 1000])
    ws2 = wb.create_sheet("Empty")
    ws2.append([None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(n_rows=3) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["t", "x"])
    for i in range(n_rows):
        w.writerow([i, i * i])
    return buf.getvalue().encode()


HTML = b"""<html><head><title>Fluid Page</title><script>var x=1;</script></head><body>
<h1>Fluids</h1><p>Water is incompressible.</p>
<h2>Table</h2><table><tr><th>Fluid</th><th>Rho</th></tr><tr><td>Water</td><td>1000</td></tr></table>
<ul><li>First item</li></ul></body></html>"""

MD = b"""# Notes

Intro paragraph.

## Data

| a | b |
|---|---|
| 1 | 2 |

Closing paragraph.
"""
